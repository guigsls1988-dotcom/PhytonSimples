from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("TruckPag_Axur_Finalizacao.pptx")

BLACK = RGBColor(8, 9, 11)
PANEL = RGBColor(20, 21, 25)
PANEL_2 = RGBColor(29, 30, 35)
WHITE = RGBColor(247, 247, 248)
MUTED = RGBColor(174, 177, 184)
RED = RGBColor(226, 38, 48)
RED_DARK = RGBColor(109, 10, 20)
GREEN = RGBColor(52, 215, 143)
AMBER = RGBColor(255, 176, 32)
BLUE = RGBColor(68, 145, 255)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def line(slide, x1, y1, x2, y2, color, width=1.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x1), Inches(y1), Inches(x2 - x1), Inches(max(y2 - y1, 0.01))
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    shape.height = Pt(width)
    return shape


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_text(slide, segments, x, y, w, h, size=20, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    for value, color, bold in segments:
        run = p.add_run()
        run.text = value
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def base_slide(prs, number, section):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK
    rect(slide, 0, 0, 0.1, 7.5, RED)
    text(slide, "NETFIVE  ×  TRUCKPAG", 0.55, 0.28, 3.4, 0.25, 10, MUTED, True)
    text(slide, section.upper(), 9.2, 0.28, 3.55, 0.25, 10, MUTED, True, align=PP_ALIGN.RIGHT)
    text(slide, f"{number:02d}", 12.2, 7.03, 0.55, 0.2, 9, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def title(slide, eyebrow, heading, sub=None):
    text(slide, eyebrow.upper(), 0.7, 0.85, 3.5, 0.25, 11, RED, True)
    text(slide, heading, 0.7, 1.18, 11.7, 0.9, 28, WHITE, True)
    if sub:
        text(slide, sub, 0.7, 2.08, 11.5, 0.55, 15, MUTED)


def badge(slide, label, x, y, color=RED, width=1.4):
    rect(slide, x, y, width, 0.34, color, radius=True)
    text(slide, label.upper(), x, y + 0.06, width, 0.17, 9, WHITE, True, align=PP_ALIGN.CENTER)


def card(slide, x, y, w, h, kicker, heading, body, accent=RED):
    rect(slide, x, y, w, h, PANEL, radius=True)
    rect(slide, x, y, 0.06, h, accent, radius=True)
    text(slide, kicker.upper(), x + 0.25, y + 0.22, w - 0.45, 0.2, 10, accent, True)
    text(slide, heading, x + 0.25, y + 0.55, w - 0.45, 0.42, 18, WHITE, True)
    text(slide, body, x + 0.25, y + 1.05, w - 0.5, h - 1.2, 12, MUTED)


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK
    rect(slide, 9.6, 0, 3.73, 7.5, RED_DARK)
    rect(slide, 10.7, 0, 2.63, 7.5, RED)
    text(slide, "netfive", 0.75, 0.6, 2.2, 0.45, 24, WHITE, True)
    text(slide, "TruckPag", 3.0, 0.6, 2.4, 0.45, 24, GREEN, True)
    text(slide, "ENCERRAMENTO DA IMPLEMENTAÇÃO", 0.78, 2.0, 6.5, 0.32, 12, RED, True)
    text(slide, "TruckPag × Axur", 0.75, 2.42, 8.3, 0.85, 38, WHITE, True)
    text(slide, "Proteção de Presença Digital", 0.78, 3.37, 7.0, 0.45, 21, MUTED)
    line(slide, 0.78, 4.35, 6.85, 4.35, RED, 3)
    text(slide, "Finalização após o Follow-up 2", 0.78, 4.65, 5.4, 0.35, 15, WHITE, True)
    text(slide, "Validações pendentes e próximos passos", 0.78, 5.08, 5.8, 0.3, 13, MUTED)
    text(slide, "JULHO · 2026", 0.78, 6.62, 2.6, 0.25, 10, MUTED, True)
    text(slide, "RESILIÊNCIA\nDIGITAL\nATIVADA.", 9.85, 2.2, 3.0, 2.3, 25, WHITE, True)


def add_status(prs):
    slide = base_slide(prs, 2, "Status")
    title(slide, "Visão executiva", "A implementação está na etapa de finalização",
          "Não houve um terceiro follow-up. O encerramento acontece após o Follow-up 2.")
    stages = [
        ("01", "Kickoff", GREEN),
        ("02", "Configuração", GREEN),
        ("03", "Follow-up 1", GREEN),
        ("04", "Follow-up 2", GREEN),
        ("05", "Finalização", RED),
    ]
    line(slide, 1.15, 4.15, 12.15, 4.15, PANEL_2, 5)
    for i, (num, label, color) in enumerate(stages):
        x = 1.0 + i * 2.72
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.72), Inches(0.85), Inches(0.85))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.color.rgb = color
        text(slide, num, x, 3.99, 0.85, 0.18, 12, BLACK if color != RED else WHITE, True,
             align=PP_ALIGN.CENTER)
        text(slide, label, x - 0.35, 4.75, 1.55, 0.45, 12, WHITE, True, align=PP_ALIGN.CENTER)
        if i < 4:
            text(slide, "CONCLUÍDO", x - 0.25, 5.27, 1.35, 0.2, 9, GREEN, True, align=PP_ALIGN.CENTER)
        else:
            badge(slide, "Estamos aqui", x - 0.18, 5.18, RED, 1.22)


def add_summary(prs):
    slide = base_slide(prs, 3, "Resumo")
    title(slide, "Resumo executivo", "O foco agora é validar e decidir",
          "A plataforma já gerou sinais úteis; o encerramento depende de decisões da TruckPag.")
    card(slide, 0.7, 2.9, 3.85, 2.65, "Ativado", "Monitoramento operacional",
         "A jornada de implantação e os dois follow-ups foram concluídos.", GREEN)
    card(slide, 4.75, 2.9, 3.85, 2.65, "Validar", "Proteção de Marca",
         "Revisar os achados detectados e confirmar a tratativa dos tickets.", AMBER)
    card(slide, 8.8, 2.9, 3.85, 2.65, "Decidir", "OnePixel",
         "Definir se a funcionalidade será implementada ou ficará fora do escopo atual.", RED)
    rect(slide, 0.7, 5.88, 11.95, 0.68, PANEL_2, radius=True)
    rich_text(slide, [
        ("RESULTADO ESPERADO  ", RED, True),
        ("encerrar a implantação com responsáveis, decisões e próximos passos registrados.", WHITE, False),
    ], 1.0, 6.1, 11.35, 0.25, 13)


def add_brand(prs):
    slide = base_slide(prs, 4, "Proteção de Marca")
    title(slide, "Achados identificados", "Proteção de Marca exige validação da TruckPag",
          "A visão compartilhada da plataforma apresentava 8 tickets abertos.")
    findings = [
        ("01", "Phishing", "Página “Gestão de Pedágio – PagPedágio”", RED),
        ("02", "Domínio similar", "Registro associado a truckpag.notion.site", AMBER),
        ("03", "Perfil falso", "Ocorrência em rede social / LinkedIn", BLUE),
    ]
    for i, (num, label, desc, color) in enumerate(findings):
        y = 2.85 + i * 1.15
        rect(slide, 0.7, y, 7.6, 0.91, PANEL, radius=True)
        badge(slide, num, 0.95, y + 0.27, color, 0.48)
        text(slide, label, 1.65, y + 0.2, 2.2, 0.25, 14, WHITE, True)
        text(slide, desc, 3.65, y + 0.22, 4.25, 0.25, 12, MUTED)
    rect(slide, 8.65, 2.85, 4.0, 3.21, RED_DARK, radius=True)
    text(slide, "VALIDAÇÃO NECESSÁRIA", 8.95, 3.18, 3.4, 0.25, 11, RED, True)
    text(slide, "Para cada ticket:", 8.95, 3.65, 3.1, 0.3, 18, WHITE, True)
    text(slide, "• confirmar legitimidade\n• definir prioridade\n• aprovar tratativa\n• indicar responsável", 8.95, 4.15, 3.1, 1.38, 14, WHITE)
    text(slide, "Contagens conforme evidências fornecidas; confirmar o status ao vivo no encerramento.",
         0.7, 6.5, 8.8, 0.28, 10, MUTED)


def add_other_surfaces(prs):
    slide = base_slide(prs, 5, "Superfícies monitoradas")
    title(slide, "Visibilidade ampliada", "Outros sinais disponíveis na plataforma",
          "Os módulos complementam a visão de risco e ajudam a priorizar respostas.")
    card(slide, 0.7, 2.85, 5.75, 2.8, "Vazamento de dados", "1 menção observada",
         "A evidência compartilhada mostra um arquivo associado ao termo “Truckpag”, com status “Novo”. "
         "Validar o conteúdo da amostra e a necessidade de resposta.", AMBER)
    card(slide, 6.7, 2.85, 5.95, 2.8, "Deep & Dark Web", "Nenhum ticket na visão apresentada",
         "A tela compartilhada não exibia resultados. Manter o monitoramento contínuo após o encerramento "
         "da implantação.", GREEN)
    text(slide, "AÇÃO RECOMENDADA", 0.7, 6.08, 2.0, 0.25, 10, RED, True)
    text(slide, "Registrar a conclusão da análise e o responsável por cada sinal.", 2.48, 6.03, 8.9, 0.32, 14, WHITE, True)


def add_onepixel(prs):
    slide = base_slide(prs, 6, "OnePixel")
    title(slide, "Decisão pendente", "OnePixel: implementar ou manter fora do escopo?",
          "A TruckPag precisa confirmar a decisão para que o projeto seja formalmente encerrado.")
    rect(slide, 0.7, 2.8, 5.75, 2.55, PANEL, radius=True)
    badge(slide, "Opção A", 1.0, 3.1, GREEN, 1.15)
    text(slide, "Implementar", 1.0, 3.62, 2.4, 0.4, 21, WHITE, True)
    text(slide, "• definir objetivo e escopo\n• validar requisitos técnicos\n• nomear responsáveis\n• acordar data de ativação",
         1.0, 4.12, 4.85, 1.0, 13, MUTED)
    rect(slide, 6.7, 2.8, 5.95, 2.55, PANEL, radius=True)
    badge(slide, "Opção B", 7.0, 3.1, RED, 1.15)
    text(slide, "Não implementar agora", 7.0, 3.62, 4.5, 0.4, 21, WHITE, True)
    text(slide, "• registrar a decisão\n• documentar a justificativa\n• manter como evolução futura\n• definir gatilho para reavaliação",
         7.0, 4.12, 5.0, 1.0, 13, MUTED)
    rect(slide, 0.7, 5.75, 11.95, 0.75, RED_DARK, radius=True)
    text(slide, "DECISÃO DA TRUCKPAG", 1.0, 5.98, 2.7, 0.25, 11, RED, True)
    text(slide, "☐ Implementar   ☐ Não implementar agora   |   Responsável: __________   Data: ___/___/___",
         3.28, 5.94, 8.9, 0.3, 13, WHITE, True)


def add_closure(prs):
    slide = base_slide(prs, 7, "Encerramento")
    title(slide, "Checklist de conclusão", "O que falta para finalizar a implementação",
          "Quatro confirmações objetivas encerram o ciclo e iniciam a operação contínua.")
    items = [
        ("01", "Validar os tickets de Proteção de Marca", "TruckPag", AMBER),
        ("02", "Definir a tratativa para cada ocorrência", "TruckPag + Netfive", RED),
        ("03", "Decidir sobre a implementação do OnePixel", "TruckPag", BLUE),
        ("04", "Registrar responsáveis e aceite de encerramento", "TruckPag + Netfive", GREEN),
    ]
    for i, (num, item, owner, color) in enumerate(items):
        y = 2.75 + i * 0.9
        rect(slide, 0.7, y, 11.95, 0.67, PANEL, radius=True)
        badge(slide, num, 0.95, y + 0.17, color, 0.48)
        text(slide, item, 1.65, y + 0.17, 7.35, 0.24, 14, WHITE, True)
        text(slide, owner.upper(), 9.1, y + 0.19, 3.05, 0.2, 10, color, True, align=PP_ALIGN.RIGHT)
    text(slide, "SAÍDA", 0.7, 6.52, 0.7, 0.22, 10, RED, True)
    text(slide, "Implantação encerrada → operação e acompanhamento contínuos", 1.45, 6.46, 8.5, 0.3, 15, WHITE, True)


def add_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK
    rect(slide, 0, 0, 0.12, 7.5, RED)
    text(slide, "ENCERRAMENTO", 0.85, 1.08, 2.2, 0.25, 11, RED, True)
    text(slide, "Decisões claras.\nProteção contínua.", 0.85, 1.62, 8.8, 1.6, 36, WHITE, True)
    text(slide, "Próximo passo: concluir as validações e registrar o aceite da TruckPag.",
         0.88, 3.72, 8.9, 0.48, 17, MUTED)
    line(slide, 0.88, 4.63, 5.95, 4.63, RED, 3)
    text(slide, "NETFIVE  ×  TRUCKPAG", 0.88, 5.05, 3.8, 0.3, 14, WHITE, True)
    text(slide, "Obrigado.", 9.3, 5.65, 3.2, 0.6, 28, RED, True, align=PP_ALIGN.RIGHT)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "TruckPag × Axur — Finalização da Implementação"
    prs.core_properties.subject = "Validações pendentes de Proteção de Marca e OnePixel"
    prs.core_properties.author = "Netfive"
    add_cover(prs)
    add_status(prs)
    add_summary(prs)
    add_brand(prs)
    add_other_surfaces(prs)
    add_onepixel(prs)
    add_closure(prs)
    add_close(prs)
    prs.save(OUT)
    print(f"Created {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
